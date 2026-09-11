import os
import re
import hashlib
import logging
import tempfile
import atexit
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    import fitz as pymupdf
except ImportError:
    pymupdf = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from paddleocr import PaddleOCR
except ImportError:
    PaddleOCR = None

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".xlsx", ".csv", ".pptx", ".png", ".jpg", ".jpeg", ".bmp", ".gif"}


class ParsedDocument:
    def __init__(self, file_path: str, content: str, metadata: dict, pages: Optional[list[str]] = None):
        self.file_path = file_path
        self.content = content
        self.metadata = metadata
        self.pages = pages or []
        self.document_id = hashlib.md5(f"{file_path}_{os.path.getmtime(file_path)}".encode()).hexdigest()


def parse_document(file_path: str, ocr_enabled: bool = False, vision_fn=None) -> Optional[ParsedDocument]:
    path = Path(file_path)
    if not path.exists():
        return None

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        return None

    metadata = {
        "filename": path.name,
        "file_type": ext.lstrip("."),
        "file_size": path.stat().st_size,
        "file_path": str(path.absolute()),
        "title": path.stem,
        "source": "local_upload",
        "author": "",
        "department": "",
        "classification": "internal",
        "page_count": 0,
    }

    content = ""
    pages = []
    try:
        if ext == ".pdf":
            content, pages = _parse_pdf(file_path, ocr_enabled)
            metadata["page_count"] = len(pages)
        elif ext in (".txt", ".md", ".csv"):
            content = _parse_text(file_path)
        elif ext == ".docx":
            content = _parse_docx(file_path)
        elif ext == ".xlsx":
            content = _parse_xlsx(file_path)
        elif ext == ".pptx":
            content, pages = _parse_pptx(file_path)
            metadata["page_count"] = len(pages)
        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".gif"):
            content = _parse_image(file_path, vision_fn)
            metadata["file_type"] = "image"
    except Exception as e:
        logger.error("Parse error for %s: %s", file_path, e)
        metadata["parse_error"] = str(e)
        content = ""

    _extract_metadata_hints(content, metadata)

    return ParsedDocument(file_path=file_path, content=content, metadata=metadata, pages=pages)


def _extract_metadata_hints(content: str, metadata: dict):
    if not metadata.get("author"):
        author_patterns = [r'作者[：:]\s*(.+?)(?:\n|$)', r'Author[：:]\s*(.+?)(?:\n|$)']
        for p in author_patterns:
            m = re.search(p, content[:500])
            if m:
                metadata["author"] = m.group(1).strip()
                break

    if not metadata.get("department"):
        dept_patterns = [r'部门[：:]\s*(.+?)(?:\n|$)', r'Department[：:]\s*(.+?)(?:\n|$)']
        for p in dept_patterns:
            m = re.search(p, content[:500])
            if m:
                metadata["department"] = m.group(1).strip()
                break

    classification_keywords = {"机密": "confidential", "绝密": "top_secret", "内部": "internal", "公开": "public"}
    for keyword, level in classification_keywords.items():
        if keyword in content[:1000]:
            metadata["classification"] = level
            break


def _parse_pdf_ocr(file_path: str) -> str:
    if pymupdf is None and PaddleOCR is None:
        return ""
    pages_text = []
    try:
        doc = pymupdf.open(file_path)
        if PaddleOCR is not None:
            ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
        else:
            ocr = None

        for i, page in enumerate(doc, 1):
            page_text = page.get_text()
            if page_text.strip() and len(page_text.strip()) > 20:
                pages_text.append(f"--- 第{i}页 ---\n{page_text}")
                continue

            if ocr is not None:
                pix = page.get_pixmap(dpi=200)
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_file:
                    img_path = tmp_file.name
                pix.save(img_path)
                try:
                    result = ocr.ocr(img_path, cls=True)
                    ocr_text = ""
                    if result and result[0]:
                        ocr_text = "\n".join(
                            line[1][0] for line in result[0] if line and len(line) > 1
                        )
                    pages_text.append(f"--- 第{i}页 (OCR) ---\n{ocr_text or page_text}")
                except Exception:
                    pages_text.append(f"--- 第{i}页 ---\n{page_text}")
                finally:
                    try:
                        os.remove(img_path)
                    except OSError:
                        pass
            else:
                pages_text.append(f"--- 第{i}页 ---\n{page_text or '[扫描件-无文本]'}")

        doc.close()
    except Exception as e:
        logger.warning("PyMuPDF/OCR failed: %s, falling back to PyPDF2", e)
        return _parse_pdf_fallback(file_path)

    return "\n\n".join(pages_text)


def _parse_pdf(file_path: str, ocr_enabled: bool = False) -> tuple[str, list[str]]:
    if ocr_enabled and pymupdf is not None:
        text = _parse_pdf_ocr(file_path)
        pages = [p for p in text.split("\n\n---") if p.strip()]
        return text, pages

    if pymupdf is not None:
        doc = pymupdf.open(file_path)
        pages = []
        for i, page in enumerate(doc, 1):
            text = page.get_text() or ""
            pages.append(f"--- 第{i}页 ---\n{text}")
        doc.close()
        return "\n\n".join(pages), pages

    return _parse_pdf_fallback(file_path), []


def _parse_pdf_fallback(file_path: str) -> str:
    if PdfReader is None:
        return _parse_text_fallback(file_path)
    reader = PdfReader(file_path)
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        pages.append(f"--- 第{i}页 ---\n{text}")
    return "\n\n".join(pages)


def _parse_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _parse_text_fallback(file_path: str) -> str:
    with open(file_path, "rb") as f:
        raw = f.read()
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="replace")


def _parse_docx(file_path: str) -> str:
    if DocxDocument is None:
        return _parse_text_fallback(file_path)
    doc = DocxDocument(file_path)
    parts = []

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()
        if not text:
            continue
        if "Heading" in style_name or "标题" in style_name:
            level = re.search(r'\d+', style_name)
            prefix = "#" * int(level.group()) if level else "##"
            parts.append(f"\n{prefix} {text}\n")
        else:
            parts.append(text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            parts.append("\n" + "\n".join(rows) + "\n")

    return "\n".join(parts)


def _parse_xlsx(file_path: str) -> str:
    if load_workbook is None:
        return _parse_text_fallback(file_path)
    wb = load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"=== 工作表: {sheet_name} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_pptx(file_path: str) -> tuple[str, list[str]]:
    if Presentation is None:
        return _parse_text_fallback(file_path), []
    prs = Presentation(file_path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        parts = [f"=== 幻灯片 {i} ==="]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        level = para.level or 0
                        prefix = "  " * level + ("- " if level > 0 else "")
                        parts.append(f"{prefix}{text}")
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    parts.append("| " + " | ".join(cells) + " |")
        slides.append("\n".join(parts))
    return "\n\n".join(slides), slides


def _parse_image(file_path: str, vision_fn=None) -> str:
    if vision_fn:
        try:
            import base64
            with open(file_path, "rb") as f:
                img_data = base64.b64encode(f.read()).decode()
            return vision_fn(img_data)
        except Exception as e:
            logger.error("Vision parsing failed: %s", e)

    metadata_lines = [
        f"[图片文件: {os.path.basename(file_path)}]",
        f"[大小: {os.path.getsize(file_path)} bytes]",
    ]
    return "\n".join(metadata_lines)